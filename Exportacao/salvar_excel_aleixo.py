import pandas as pd
from openpyxl import load_workbook


def validar_totais(totais_pdf, totais_calculados):

    resultado = {}

    for verba, valor_pdf in totais_pdf.items():

        valor_calculado = totais_calculados.get(verba)

        if valor_calculado is None:
            resultado[verba] = False
            continue

        resultado[verba] = round(valor_pdf, 2) == round(valor_calculado, 2)

    return resultado


def salvar_excel_aleixo(df, caminho_saida, totais_pdf=None):

    if totais_pdf is not None:

        totais_calculados = {}

        for verba in totais_pdf:

            # Regra especifica do Aleixo
            if verba == "Cota Extra":

                total_coluna_agua = (
                    df["Taxa Extra Coluna de Agua"]
                    .fillna(0)
                    .sum()
                )

                total_laje = (
                    df["Cota Extra Laje"]
                    .fillna(0)
                    .sum()
                )

                totais_calculados["Cota Extra"] = (
                    total_coluna_agua + total_laje
                )

            elif verba in df.columns:

                totais_calculados[verba] = (
                    df[verba]
                    .fillna(0)
                    .sum()
                )

        print("\n=== Totais Calculados ===")
        print(totais_calculados)

        resultado_validacao = validar_totais(
            totais_pdf,
            totais_calculados
        )

        print("\n=== Validação ===")

        for verba, ok in resultado_validacao.items():
            print(f"{verba}: {'OK' if ok else 'ERRO'}")

    # Substitui valores 0 por vazio
    df = df.replace(0, None)

    # Preenche valores nulos com células vazias
    df = df.fillna("")

    # Remove colunas totalmente vazias
    df = df.loc[:, (df != "").any()]

    # Exporta para Excel
    df.to_excel(caminho_saida, index=False)

    if totais_pdf is not None:

        workbook = load_workbook(caminho_saida)
        planilha = workbook.active

        linha = len(df) + 3

        planilha.cell(row=linha, column=1).value = "Verba"
        planilha.cell(row=linha, column=2).value = "PDF"
        planilha.cell(row=linha, column=3).value = "Calculado"
        planilha.cell(row=linha, column=4).value = "Status"

        linha += 1

        for verba in totais_pdf:

            valor_pdf = round(
                totais_pdf.get(verba, 0),
                2
            )

            valor_calculado = round(
                totais_calculados.get(verba, 0),
                2
            )

            if valor_pdf == 0 and valor_calculado == 0:
                continue

            planilha.cell(
                row=linha,
                column=1
            ).value = verba

            planilha.cell(
                row=linha,
                column=2
            ).value = valor_pdf

            planilha.cell(
                row=linha,
                column=3
            ).value = valor_calculado

            if resultado_validacao[verba]:
                planilha.cell(
                    row=linha,
                    column=4
                ).value = "OK"
            else:
                planilha.cell(
                    row=linha,
                    column=4
                ).value = "ERRO"

            linha += 1

        linha += 1

        if all(resultado_validacao.values()):

            planilha.cell(
                row=linha,
                column=1
            ).value = "VALIDAÇÃO GERAL"

            planilha.cell(
                row=linha,
                column=2
            ).value = "APROVADA"

        else:

            planilha.cell(
                row=linha,
                column=1
            ).value = "VALIDAÇÃO GERAL"

            planilha.cell(
                row=linha,
                column=2
            ).value = "REPROVADA"

        workbook.save(caminho_saida)